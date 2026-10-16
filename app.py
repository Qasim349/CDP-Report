import io, math, os, shutil, subprocess, tempfile, zipfile
from pathlib import Path
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from flask import Flask, render_template_string, request, send_file
from openpyxl import load_workbook
from docx import Document
from docx.shared import Inches, Emu

app = Flask(__name__)
app.config["MAX_CONTENT_LENGTH"] = 25 * 1024 * 1024
ROOT = Path(__file__).resolve().parent
TEMPLATE = ROOT / "report_template.docx"
INPUT = ROOT / "report_input_template.xlsx"

HTML = '''
<!doctype html><html><head><meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1">
<title>QC Report Generator</title>
<style>
body{font-family:Arial,sans-serif;max-width:850px;margin:35px auto;padding:0 18px;color:#222}
.panel{border:1px solid #ddd;border-radius:12px;padding:22px;margin:18px 0}
button,.button{display:inline-block;background:#1f5fbf;color:white;border:0;border-radius:8px;padding:12px 18px;text-decoration:none;cursor:pointer}
label{display:block;font-weight:bold;margin:16px 0 7px}input[type=file]{width:100%;padding:10px;box-sizing:border-box}
.muted{color:#555}.error{background:#fff1f0;padding:12px;border-radius:8px}
</style></head><body>
<h1>QC Comparative Dissolution Report Generator</h1>
<p class="muted">Upload your completed Excel workbook to create a Word report and PDF using your supplied report layout.</p>
<div class="panel"><h2>1. Download Excel input template</h2>
<p>Includes Product, pH 1.2, pH 4.5 and pH 6.8 sheets. Example values are from the uploaded report; replace them for each new study.</p>
<a class="button" href="/download-input">Download Excel template</a></div>
<div class="panel"><h2>2. Generate report</h2>
{% if error %}<p class="error"><strong>Could not generate:</strong> {{error}}</p>{% endif %}
<form method="post" enctype="multipart/form-data">
<label>Completed Excel workbook (.xlsx)</label><input name="input_file" type="file" accept=".xlsx" required>
<p class="muted">Output is a ZIP file containing the editable Word report and PDF.</p><button type="submit">Generate Word + PDF</button>
</form></div>
<p class="muted">Review calculations and output under your approved QC procedure before official use.</p>
</body></html>
'''

SHEETS = [("pH 1.2", "1.2"), ("pH 4.5", "4.5"), ("pH 6.8", "6.8")]

def read_rows(sheet):
    rows = []
    for time, test, ref in sheet.iter_rows(min_row=2, max_col=3, values_only=True):
        if time is None and test is None and ref is None:
            continue
        if time is None or test is None or ref is None:
            raise ValueError(f"{sheet.title}: every row needs Time, Test, and Reference values.")
        try:
            time, test, ref = float(time), float(test), float(ref)
        except (TypeError, ValueError):
            raise ValueError(f"{sheet.title}: data values must be numeric.")
        if not 0 <= test <= 100 or not 0 <= ref <= 100:
            raise ValueError(f"{sheet.title}: Test and Reference must be from 0 to 100.")
        rows.append((time, test, ref))
    if not rows:
        raise ValueError(f"{sheet.title} has no data rows.")
    if len(rows) > 30:
        raise ValueError(f"{sheet.title} supports a maximum of 30 time points.")
    return rows

def f2_result(rows):
    squares = [(r - t) ** 2 for _, t, r in rows]
    total = sum(squares)
    value = 50 * math.log10(100 / math.sqrt(1 + total / len(rows)))
    return squares, total, value

def set_text(cell, value):
    cell.text = str(value)

def set_paragraph(p, text):
    if p.runs:
        p.runs[0].text = text
        for run in p.runs[1:]:
            run.text = ""
    else:
        p.add_run(text)

def make_chart(rows, ph, path):
    times = [x[0] for x in rows]
    tests = [x[1] for x in rows]
    refs = [x[2] for x in rows]
    fig, ax = plt.subplots(figsize=(8.4, 3.75))
    ax.plot(times, tests, marker="o", linewidth=2, label="Test")
    ax.plot(times, refs, marker="o", linewidth=2, label="Reference (Innovator)")
    ax.set_title(f"% Valsartan at pH {ph}")
    ax.set_xlabel("Time (minutes)")
    ax.set_ylabel("Dissolution (%)")
    ax.set_ylim(0, max(100, math.ceil(max(tests + refs) / 10) * 10))
    ax.grid(True, alpha=0.25)
    ax.legend()
    fig.tight_layout()
    fig.savefig(path, dpi=180, bbox_inches="tight")
    plt.close(fig)

def replace_charts(doc, paths):
    chart_paras = [p for p in doc.paragraphs if p._p.xpath(".//w:drawing")]
    if len(chart_paras) < 3:
        raise ValueError("Could not find the three charts in the Word template.")
    for p, path in zip(chart_paras[:3], paths):
        ext = p._p.xpath(".//wp:extent")
        width, height = Inches(6.75), Inches(3.0)
        if ext:
            try:
                width, height = Emu(int(ext[0].get("cx"))), Emu(int(ext[0].get("cy")))
            except Exception:
                pass
        for drawing in list(p._p.xpath(".//w:drawing")):
            drawing.getparent().remove(drawing)
        p.add_run().add_picture(path, width=width, height=height)

def generate(xlsx_path, work):
    wb = load_workbook(xlsx_path, data_only=True)
    if "Product" not in wb.sheetnames:
        raise ValueError("Missing Product sheet. Use the downloadable Excel input template.")
    pws = wb["Product"]
    details = {str(pws.cell(r, 1).value).strip(): pws.cell(r, 2).value
               for r in range(2, pws.max_row + 1) if pws.cell(r, 1).value}
    product_name = str(details.get("Product name") or "Product")
    units = int(details.get("Units evaluated") or 12)
    all_results = []
    for sheet_name, ph in SHEETS:
        if sheet_name not in wb.sheetnames:
            raise ValueError(f"Missing sheet: {sheet_name}")
        rows = read_rows(wb[sheet_name])
        squares, total, f2 = f2_result(rows)
        all_results.append((sheet_name, ph, rows, squares, total, f2))

    doc = Document(str(TEMPLATE))
    tables = doc.tables
    data_indexes, sum_indexes, result_indexes = [0, 3, 6], [1, 4, 7], [2, 5, 8]
    for i, (sheet_name, ph, rows, squares, total, f2) in enumerate(all_results):
        table = tables[data_indexes[i]]
        while len(table.rows) < len(rows) + 1:
            table.add_row()
        while len(table.rows) > len(rows) + 1:
            tr = table.rows[-1]._tr
            tr.getparent().remove(tr)
        headers = ["Time (mins.)", "Test Tt", "Reference Rt", "Reference - Test", "Squared difference"]
        for c, heading in enumerate(headers):
            set_text(table.cell(0, c), heading)
        for ri, (time, test, ref) in enumerate(rows, start=1):
            diff = ref - test
            vals = [f"{time:g}", f"{test:.2f}", f"{ref:.2f}", f"{diff:.2f}", f"{diff*diff:.4f}"]
            for ci, val in enumerate(vals):
                set_text(table.cell(ri, ci), val)
        set_text(tables[sum_indexes[i]].cell(0, 0), f"{total:.2f}")
        set_text(tables[result_indexes[i]].cell(0, 0), f"f2 = {f2:.2f}")

        for p in doc.paragraphs:
            if "STATISTICAL EVALUATION OF" in p.text and f"pH {ph}" in p.text:
                set_paragraph(p, f"STATISTICAL EVALUATION OF {units} UNITS AT pH {ph}")
            if p.text.strip() == "Valsartan (10/160 mg Tablets)":
                set_paragraph(p, product_name)

    # Refresh the six explanatory calculation lines after each "Where:" heading.
    paragraphs = doc.paragraphs
    where_indexes = [i for i, p in enumerate(paragraphs) if p.text.strip() == "Where:"]
    for section_i, (_, ph, rows, squares, total, f2) in enumerate(all_results):
        if section_i >= len(where_indexes):
            break
        pi = where_indexes[section_i]
        mean_sq = total / len(rows)
        factor = 1 / math.sqrt(1 + mean_sq)
        lines = [
            f"f2 = 50 log10 {{[1 + (1/{len(rows)}) x ({total:.2f})]^-0.5 x 100}}",
            f"f2 = 50 log10 {{[1 + ({mean_sq:.4f})]^-0.5 x 100}}",
            f"f2 = 50 log10 {{[{1+mean_sq:.4f}]^-0.5 x 100}}",
            f"f2 = 50 log10 {{{factor:.4f} x 100}}",
            f"f2 = 50 log10 {{{factor*100:.2f}}}",
            f"f2 = 50 x {math.log10(factor*100):.4f}",
        ]
        for j, line in enumerate(lines, start=2):
            if pi + j < len(paragraphs):
                set_paragraph(paragraphs[pi+j], line)

    chart_paths = []
    for i, (_, ph, rows, _, _, _) in enumerate(all_results):
        path = os.path.join(work, f"chart_{i+1}.png")
        make_chart(rows, ph, path)
        chart_paths.append(path)
    replace_charts(doc, chart_paths)
    out_docx = os.path.join(work, "generated_report.docx")
    doc.save(out_docx)
    return out_docx

def to_pdf(docx_path, work):
    office = shutil.which("libreoffice") or shutil.which("soffice")
    if not office:
        raise RuntimeError("LibreOffice is unavailable. Check Render's apt.txt setup.")
    profile = Path(work) / "lo_profile"
    result = subprocess.run([office, f"-env:UserInstallation=file://{profile}", "--headless",
        "--convert-to", "pdf", "--outdir", work, docx_path], capture_output=True, text=True, timeout=120)
    pdf = os.path.join(work, Path(docx_path).stem + ".pdf")
    if not os.path.exists(pdf):
        raise RuntimeError("PDF conversion failed: " + (result.stderr or result.stdout or "Unknown error"))
    return pdf

@app.route("/", methods=["GET", "POST"])
def index():
    if request.method == "GET":
        return render_template_string(HTML, error=None)
    upload = request.files.get("input_file")
    if not upload or not upload.filename.lower().endswith(".xlsx"):
        return render_template_string(HTML, error="Upload an .xlsx workbook from the supplied input template.")
    try:
        with tempfile.TemporaryDirectory() as work:
            xlsx = os.path.join(work, "input.xlsx")
            upload.save(xlsx)
            docx = generate(xlsx, work)
            pdf = to_pdf(docx, work)
            archive = os.path.join(work, "reports.zip")
            with zipfile.ZipFile(archive, "w", zipfile.ZIP_DEFLATED) as z:
                z.write(docx, "generated_report.docx")
                z.write(pdf, "generated_report.pdf")
            return send_file(io.BytesIO(Path(archive).read_bytes()), as_attachment=True,
                download_name="generated_reports.zip", mimetype="application/zip")
    except Exception as e:
        return render_template_string(HTML, error=str(e))

@app.route("/download-input")
def download_input():
    return send_file(str(INPUT), as_attachment=True, download_name="report_input_template.xlsx")

@app.get("/health")
def health():
    return {"status": "ok"}

if __name__ == "__main__":
    app.run(host="0.0.0.0", port=int(os.environ.get("PORT", 5000)))
